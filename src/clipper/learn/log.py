"""The performance log: one spreadsheet for every clip, across runs and campaigns.

Each run appends a row per accepted clip with the identifying columns filled
in; the user fills in the numbers for the ones they post, reading them off
TikTok Studio. One file, not one per source: per-source templates were written
once and never refreshed, so after a re-run they listed clips that no longer
existed.

It is an Excel file (data/performance.xlsx) rather than CSV so it can keep
column widths and a frozen header, and it leads with the caption -- the text
TikTok Studio lists each post by -- so rows can be matched without file names.
Only rows with some result logged are analysed; unposted clips stay blank.
"""

from __future__ import annotations

import csv
import json
import re
from dataclasses import asdict, dataclass
from datetime import date, datetime
from pathlib import Path

from ..paths import data_root, ensure
from ..utils.logging import get_logger

log = get_logger(__name__)

#: Column order as the user sees it: caption to match on, then the numbers in
#: the order TikTok Studio shows them, then what the tool needs to join back.
RESULT_COLUMNS = ["url", "posted_at", "views_24h", "views_7d", "views_30d", "views_latest",
                  "likes", "comments", "shares", "saves",
                  "avg_watch_s", "watched_full_pct", "drop_off_s", "new_followers",
                  "verified_views", "payout_usd", "notes"]
ID_COLUMNS = ["platform", "account", "duration_s", "opening", "lead_in_s", "hook",
              "campaign", "source_title", "file",
              "source_id", "clip_id", "candidate_id", "video_id", "synced_at", "studio_at"]
COLUMNS = ["caption", *RESULT_COLUMNS, *ID_COLUMNS]
NUMERIC = ["views_24h", "views_7d", "views_30d", "views_latest", "likes", "comments", "shares",
           "saves", "avg_watch_s", "watched_full_pct", "drop_off_s", "new_followers",
           "verified_views", "payout_usd"]
#: Filled by `clipper tiktok collect` from a copied TikTok Studio page.
COLLECTED = {"avg_watch_s", "watched_full_pct", "drop_off_s", "saves", "new_followers"}
#: Filled by `clipper tiktok sync`; the rest of the yellow columns are the user's.
SYNCED = {"url", "posted_at", "views_24h", "views_7d", "views_30d", "views_latest",
          "likes", "comments", "shares"}

#: Hover notes on the header cells.
HINTS = {
    "caption": "The caption as posted. Find the post in TikTok Studio by this text.",
    "views_24h": "Video views about a day after posting.",
    "views_7d": "Video views about a week after posting.",
    "views_latest": "Views at the last `clipper tiktok sync` (see synced_at).",
    "avg_watch_s": "Average watch time, in seconds (e.g. 9.8s). The most useful number: "
                   "it shows whether people who saw the clip stayed.",
    "watched_full_pct": "Watched full video, as a percentage (e.g. 12%).",
    "saves": "Saves / favorites.",
    "drop_off_s": "Where most viewers stopped watching, in seconds (TikTok Studio: "
                  "'Most viewers stopped watching at 0:01').",
    "new_followers": "New followers from this post.",
    "verified_views": "Views the campaign counted (e.g. on Vyro).",
    "opening": "How the clip's first shot was framed (filled in by the tool).",
    "lead_in_s": "Seconds of silence before the first word (filled in by the tool).",
    "hook": "The hook line shown on screen at the start, if any (filled in by the tool).",
    "payout_usd": "What the campaign paid for this post.",
}
_WIDE = {"caption": 60, "notes": 40, "url": 40}


def log_path() -> Path:
    return data_root() / "performance.xlsx"


@dataclass(frozen=True)
class NewClip:
    """What a run knows about a clip when it adds it to the log."""

    source_id: str
    clip_id: str
    candidate_id: str
    campaign: str
    source_title: str
    file: str
    duration_s: float
    caption: str = ""
    # How the clip opened, for comparing posts with and without each change:
    # "full-screen", "fit 4:5", "fit wide" or "letterbox"; seconds of silence
    # before the first word; the hook line shown on screen ("" for none).
    opening: str = ""
    lead_in_s: float | None = None
    hook: str = ""


def read(path: Path | None = None) -> list[dict[str, str]]:
    path = path or log_path()
    if not path.exists():
        legacy = path.with_suffix(".csv")
        if path.suffix == ".xlsx" and legacy.exists():
            return _read_csv(legacy)  # an older log; the next write converts it
        return []
    if path.suffix.lower() == ".csv":
        return _read_csv(path)
    return _read_xlsx(path)


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:  # -sig: Excel adds a BOM
        return [{k: v or "" for k, v in row.items() if k} for row in csv.DictReader(handle)]


def _read_xlsx(path: Path) -> list[dict[str, str]]:
    from openpyxl import load_workbook

    ws = load_workbook(path, data_only=True).active
    rows = ws.iter_rows()
    header = [str(c.value).strip() if c.value is not None else "" for c in next(rows, [])]
    out = []
    for cells in rows:
        row = {name: _cell_text(cell) for name, cell in zip(header, cells, strict=False) if name}
        if any(row.values()):
            out.append(row)
    return out


def _cell_text(cell) -> str:
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (int, float)):
        if "%" in (cell.number_format or ""):
            value = value * 100  # Excel stores a typed "12%" as 0.12
        return str(int(value)) if float(value).is_integer() else f"{value:g}"
    return str(value).strip()


def write(rows: list[dict[str, str]], path: Path | None = None) -> Path:
    path = path or log_path()
    ensure(path.parent)
    extra = sorted({k for r in rows for k in r} - set(COLUMNS) - {""})
    columns = COLUMNS + extra
    if path.suffix.lower() == ".csv":
        with path.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="ignore")
            writer.writeheader()
            for row in rows:
                writer.writerow({k: row.get(k, "") or "" for k in columns})
        return path
    _write_xlsx(rows, columns, path)
    legacy = path.with_suffix(".csv")
    if legacy.exists():  # converted: keep the old file, out of the way
        legacy.replace(legacy.with_suffix(".csv.old"))
    return path


def _write_xlsx(rows: list[dict[str, str]], columns: list[str], path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Posts"
    ws.append(columns)
    for row in rows:
        values = []
        for name in columns:
            text = row.get(name, "") or ""
            n = number(text) if name in NUMERIC or name == "duration_s" else None
            values.append(n if n is not None and text else text)
        ws.append(values)

    fill = PatternFill("solid", fgColor="FFF2CC")      # typed in by the user
    synced = PatternFill("solid", fgColor="DDEBF7")    # filled by `clipper tiktok sync`
    collected = PatternFill("solid", fgColor="E2EFDA")  # filled by `clipper tiktok collect`
    for i, name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True)
        if name in SYNCED:
            cell.fill = synced
        elif name in COLLECTED:
            cell.fill = collected
        elif name in NUMERIC or name == "notes":
            cell.fill = fill
        if name in HINTS:
            cell.comment = Comment(HINTS[name], "clipper")
        longest = max([len(name)] + [len(str(row.get(name, "") or "")) for row in rows])
        width = min(longest + 2, _WIDE.get(name, 30))
        ws.column_dimensions[get_column_letter(i)].width = max(width, 8)
    wrap = Alignment(wrap_text=True, vertical="top")
    top = Alignment(vertical="top")
    for cells in ws.iter_rows(min_row=2):
        for cell in cells:
            cell.alignment = wrap if columns[cell.column - 1] in _WIDE else top
    ws.freeze_panes = "B2"  # header row and caption column stay in view
    tmp = path.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    tmp.replace(path)


def add_clips(clips: list[NewClip], path: Path | None = None) -> Path:
    """Append rows for clips not in the log yet. Existing rows are never changed.

    While the spreadsheet is open in Excel, Windows refuses to replace it. The
    clips are then parked in a side file and added by the next successful
    call, rather than failing the run that produced them.
    """
    path = path or log_path()
    pending = path.with_suffix(".pending.json")
    waiting = _load_pending(pending)
    try:
        written = _add(waiting + list(clips), path)
    except PermissionError:
        _save_pending(pending, waiting + list(clips))
        log.warning("%s is open in another program; %d clip(s) will be added to it "
                    "next time", path.name, len(clips))
        return path
    pending.unlink(missing_ok=True)
    return written


def _load_pending(path: Path) -> list[NewClip]:
    try:
        return [NewClip(**item) for item in json.loads(path.read_text(encoding="utf-8"))]
    except (OSError, ValueError, TypeError):
        return []


def _save_pending(path: Path, clips: list[NewClip]) -> None:
    ensure(path.parent)
    path.write_text(json.dumps([asdict(c) for c in clips], indent=1), encoding="utf-8")


def _add(clips: list[NewClip], path: Path) -> Path:
    rows = read(path)
    known = {(r.get("source_id", ""), r.get("clip_id", "")) for r in rows}
    for clip in clips:
        if (clip.source_id, clip.clip_id) in known:
            continue
        rows.append({"caption": clip.caption, "source_id": clip.source_id,
                     "clip_id": clip.clip_id, "candidate_id": clip.candidate_id,
                     "campaign": clip.campaign, "source_title": clip.source_title,
                     "file": clip.file, "duration_s": f"{clip.duration_s:.1f}",
                     "opening": clip.opening, "hook": clip.hook,
                     "lead_in_s": "" if clip.lead_in_s is None else f"{clip.lead_in_s:.2f}"})
        known.add((clip.source_id, clip.clip_id))
    return write(rows, path)


_MIN_SEC = re.compile(r"^(?:(\d+)\s*m(?:in)?)?\s*(?:([\d.]+)\s*s(?:ec)?)?$")


def number(value: str | None) -> float | None:
    """A logged number, tolerating what people type or paste from TikTok Studio:
    '1,234', '12.5%', '3.4K', '1.2M', '$12.50', '9.8s', '1m 5s', '0:09'."""
    if value is None:
        return None
    text = str(value).strip().replace(",", "").replace("%", "").replace("$", "")
    if not text:
        return None
    if ":" in text:  # m:ss
        mins, _, secs = text.partition(":")
        try:
            return int(mins) * 60 + float(secs)
        except ValueError:
            return None
    if text[-1:].lower() == "s" or "m " in text.lower():
        match = _MIN_SEC.match(text.lower())
        if match and any(match.groups()):
            return int(match.group(1) or 0) * 60 + float(match.group(2) or 0)
        return None
    scale = 1.0
    if text[-1:].lower() in ("k", "m"):
        scale = 1_000.0 if text[-1].lower() == "k" else 1_000_000.0
        text = text[:-1]
    try:
        return float(text) * scale
    except ValueError:
        return None


def has_results(row: dict[str, str]) -> bool:
    return any(number(row.get(k)) is not None for k in NUMERIC)
