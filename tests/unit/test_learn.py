"""`clipper learn` and the performance log, on synthetic results.

BUILD_BRIEF.md 14.2: under about 20 rows, describe and change nothing; with
enough, propose weights with shrinkage toward the current ones, and write only
with --apply.
"""

from __future__ import annotations

import json

import pytest

from clipper.learn import analyze
from clipper.learn import log as perf

WEIGHTS = {"llm": 0.5, "heatmap": 0.2, "audio": 0.15, "text": 0.15}


def make_source(root, sid, n, *, audio_of=lambda i: i / 30, llm_of=lambda i: 0.5):
    """Score artifacts for `n` candidates in data/work/<sid>."""
    work = root / "work" / sid
    work.mkdir(parents=True)
    scored = [{"candidate_id": f"c{i}", "composite": 0.5 + i / 100,
               "components": {"llm": llm_of(i), "audio": audio_of(i), "text": 0.5}}
              for i in range(n)]
    signals = [{"candidate_id": f"c{i}", "llm_total": 5.0,
                "llm_a": {"hook_strength": 5 + i % 3, "payoff": 6, "hook_text": f"hook {i}"}}
               for i in range(n)]
    cands = [{"candidate_id": f"c{i}", "text": f"transcript {i}"} for i in range(n)]
    (work / "scored.json").write_text(json.dumps({"scored": scored}), encoding="utf-8")
    (work / "signals.json").write_text(json.dumps({"values": signals}), encoding="utf-8")
    (work / "candidates.json").write_text(json.dumps({"candidates": cands}), encoding="utf-8")


def rows_for(sid, n, *, watch_of=lambda i: 10.0 + i, views_of=lambda i: 300 + 10 * i):
    return [{"source_id": sid, "clip_id": f"{i:03d}", "candidate_id": f"c{i}",
             "duration_s": "40", "avg_watch_s": str(watch_of(i)),
             "views_24h": str(views_of(i)), "likes": "10"} for i in range(n)]


class TestNumbers:
    @pytest.mark.parametrize("typed, value", [
        ("1,234", 1234.0), ("45%", 45.0), ("1.2K", 1200.0), ("3M", 3_000_000.0),
        ("$12.50", 12.5), ("9.8s", 9.8), ("1m 5s", 65.0), ("0:09", 9.0),
        ("", None), ("n/a", None), (None, None)])
    def test_what_people_type_is_read(self, typed, value):
        assert perf.number(typed) == value

    def test_a_row_without_numbers_is_not_a_result(self):
        assert not perf.has_results({"clip_id": "001", "url": "https://tiktok.com/x"})
        assert perf.has_results({"clip_id": "001", "views_24h": "312"})


class TestLog:
    def test_clips_are_added_once_and_logged_numbers_survive(self, data_root):
        clip = perf.NewClip("s1", "001", "c0", "FX", "Ep 201", "001.mp4", 46.3)
        perf.add_clips([clip])
        rows = perf.read()
        rows[0]["views_24h"] = "500"
        perf.write(rows)
        perf.add_clips([clip, perf.NewClip("s1", "002", "c1", "FX", "Ep 201", "002.mp4", 33.0)])
        rows = perf.read()
        assert [(r["clip_id"], r["views_24h"]) for r in rows] == [("001", "500"), ("002", "")]
        assert rows[0]["duration_s"] == "46.3"

    def test_an_older_csv_log_is_converted_keeping_its_numbers(self, data_root):
        legacy = data_root / "performance.csv"
        legacy.write_text("﻿source_id,clip_id,candidate_id,views_24h\ns0,000,c9,777\n",
                          encoding="utf-8")  # with the byte-order mark Excel adds
        perf.add_clips([perf.NewClip("s1", "001", "c0", "FX", "", "f", 30)])
        rows = perf.read()
        assert [(r["source_id"], r["views_24h"]) for r in rows] == [("s0", "777"), ("s1", "")]
        assert not legacy.exists() and (data_root / "performance.csv.old").exists()

    def test_clips_wait_while_the_log_is_open_in_excel(self, data_root, monkeypatch):
        perf.add_clips([perf.NewClip("s1", "001", "c0", "FX", "", "f", 30)])
        real = perf._write_xlsx

        def locked(*a, **k):
            raise PermissionError("in use")
        monkeypatch.setattr(perf, "_write_xlsx", locked)
        perf.add_clips([perf.NewClip("s1", "002", "c1", "FX", "", "f", 30)])  # no crash
        assert [r["clip_id"] for r in perf.read()] == ["001"]
        monkeypatch.setattr(perf, "_write_xlsx", real)
        perf.add_clips([perf.NewClip("s1", "003", "c2", "FX", "", "f", 30)])
        assert [r["clip_id"] for r in perf.read()] == ["001", "002", "003"]
        assert not (data_root / "performance.pending.json").exists()

    def test_the_spreadsheet_is_readable_at_a_glance(self, data_root):
        from openpyxl import load_workbook

        perf.add_clips([perf.NewClip("s1", "001", "c0", "FX", "", "f", 30,
                                     caption="A long caption " * 8)])
        ws = load_workbook(perf.log_path()).active
        assert ws.freeze_panes == "B2", "header and caption stay in view"
        assert ws.column_dimensions["A"].width == 60, "long captions wrap at a fixed width"
        assert ws.cell(1, 1).value == "caption"
        assert ws["A1"].font.bold

    def test_a_percentage_typed_in_excel_reads_as_percent(self, data_root):
        from openpyxl import load_workbook

        perf.add_clips([perf.NewClip("s1", "001", "c0", "FX", "", "f", 30)])
        wb = load_workbook(perf.log_path())
        ws = wb.active
        col = perf.COLUMNS.index("watched_full_pct") + 1
        ws.cell(2, col).value = 0.12   # what Excel stores when you type 12%
        ws.cell(2, col).number_format = "0%"
        ws.cell(2, perf.COLUMNS.index("avg_watch_s") + 1).value = "9.8s"
        wb.save(perf.log_path())
        (row,) = perf.read()
        assert perf.number(row["watched_full_pct"]) == pytest.approx(12.0)
        assert perf.number(row["avg_watch_s"]) == pytest.approx(9.8)


class TestFewRows:
    def test_under_twenty_rows_describes_and_proposes_nothing(self, data_root):
        make_source(data_root, "s1", 7)
        result = analyze.analyse(rows_for("s1", 7), WEIGHTS)
        assert len(result.joined) == 7
        assert result.proposal is None and result.correlations == []
        assert "20 are needed" in result.note

    def test_watch_through_is_watch_time_over_length(self, data_root):
        make_source(data_root, "s1", 1)
        (j,) = analyze.analyse(rows_for("s1", 1, watch_of=lambda i: 20.0), WEIGHTS).joined
        assert j.watch_through == pytest.approx(0.5)
        assert j.scores["hook_strength"] == 5 and j.hook == "hook 0"

    def test_rows_without_scores_are_reported_not_guessed(self, data_root):
        make_source(data_root, "s1", 2)
        rows = [*rows_for("s1", 2), {"source_id": "gone", "clip_id": "009",
                                     "candidate_id": "c9", "views_24h": "10"}]
        result = analyze.analyse(rows, WEIGHTS)
        assert len(result.joined) == 2 and len(result.unmatched) == 1


class TestEnoughRows:
    def test_a_signal_that_predicts_watch_through_gains_weight_by_a_small_step(self, data_root):
        # audio rises with watch time; llm is noise.
        make_source(data_root, "s1", 24, llm_of=lambda i: [0.2, 0.9, 0.5, 0.1][i % 4])
        result = analyze.analyse(rows_for("s1", 24), WEIGHTS)
        audio = next(c for c in result.correlations
                     if c.name == "audio" and c.target == "watch_through")
        assert audio.rho > 0.9 and audio.low > 0.5
        p = result.proposal
        assert p is not None
        assert p.proposed["audio"] > WEIGHTS["audio"]
        assert p.proposed["audio"] - WEIGHTS["audio"] < 0.2, "moves only part of the way"
        assert p.proposed["heatmap"] == WEIGHTS["heatmap"], "absent signals are untouched"

    def test_no_positive_evidence_proposes_nothing(self, data_root):
        make_source(data_root, "s1", 24, audio_of=lambda i: 1 - i / 30,
                    llm_of=lambda i: 1 - i / 30)
        # text is constant (rho 0); llm and audio fall as watch time rises.
        assert analyze.analyse(rows_for("s1", 24), WEIGHTS).proposal is None

    def test_apply_rewrites_only_the_weights_and_keeps_comments(self, data_root, tmp_path):
        config = tmp_path / "default.yaml"
        config.write_text("weights:                     # renormalized\n  llm: 0.50\n"
                          "  heatmap: 0.20\n  audio: 0.15\n  text: 0.15\n\n"
                          "selection:\n  top_n: 5\n  audio: 0.99\n", encoding="utf-8")
        proposal = analyze.Proposal(current=WEIGHTS, proposed={**WEIGHTS, "audio": 0.2,
                                                               "llm": 0.45}, evidence=[])
        history = tmp_path / "weights_history.json"
        analyze.apply(proposal, config, n=24, history_path=history)
        text = config.read_text(encoding="utf-8")
        assert "  llm: 0.450\n" in text and "  audio: 0.200\n" in text
        assert "# renormalized" in text
        assert "  audio: 0.99" in text, "keys outside weights: are left alone"
        (entry,) = json.loads(history.read_text(encoding="utf-8"))
        assert entry["clips"] == 24 and entry["to"]["audio"] == 0.2

    def test_examples_use_the_keys_the_scoring_prompt_reads(self, data_root):
        make_source(data_root, "s1", 24)
        result = analyze.analyse(rows_for("s1", 24), WEIGHTS)
        path = analyze.write_examples(result.joined)
        first = json.loads(path.read_text(encoding="utf-8").splitlines()[0])
        assert set(first) >= {"text", "hook", "views"}
        assert first["text"] == "transcript 23", "best watch-through first"
        assert isinstance(first["views"], int)


class TestRerender:
    def test_an_unposted_rows_description_follows_a_rerender(self, data_root):
        perf.add_clips([perf.NewClip("s1", "001", "c0", "x", "", "a.mp4", 30, opening="letterbox")])
        perf.add_clips([perf.NewClip("s1", "001", "c0", "x", "", "a.mp4", 31, opening="full-screen")])
        (row,) = perf.read()
        assert row["opening"] == "full-screen" and perf.number(row["duration_s"]) == 31

    def test_a_posted_row_is_left_alone(self, data_root):
        perf.add_clips([perf.NewClip("s1", "001", "c0", "x", "", "a.mp4", 30, opening="letterbox")])
        rows = perf.read()
        rows[0]["views_24h"] = "500"
        perf.write(rows)
        perf.add_clips([perf.NewClip("s1", "001", "c0", "x", "", "a.mp4", 31, opening="full-screen")])
        (row,) = perf.read()
        assert row["opening"] == "letterbox" and row["views_24h"] == "500"
